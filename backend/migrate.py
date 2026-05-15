"""
One-time migration: reads Physical Verification of Fixed Assets.xlsx
and populates the PostgreSQL database.

Run from the backend/ directory:
    python migrate.py --excel "../Physical Verification of Fixed Assets.xlsx"

Flags:
    --excel PATH   path to the .xlsx file  (default: ../Physical Verification of Fixed Assets.xlsx)
    --clear        drop + re-insert all rows (safe to re-run)
"""

import argparse
import os
import sys
from datetime import datetime

from openpyxl import load_workbook
from sqlalchemy.orm import Session

# Make app/ importable
sys.path.insert(0, os.path.dirname(__file__))

from app.auth.jwt import hash_password
from app.database import Base, SessionLocal, engine
from app.models.models import (
    AppState, Asset, AssetClassCount, Barcode, Class, EditLog, Stickering, User,
)

# ── helpers ──────────────────────────────────────────────────────────────────

def ws_rows(wb, name: str):
    ws = wb[name] if name in wb.sheetnames else None
    if not ws:
        print(f"  [skip] sheet '{name}' not found")
        return []
    return list(ws.iter_rows(values_only=True))


def str_or_empty(v) -> str:
    return str(v).strip() if v is not None else ""


# ── main ─────────────────────────────────────────────────────────────────────

def migrate(excel_path: str, clear: bool) -> None:
    print(f"Loading workbook: {excel_path}")
    wb = load_workbook(excel_path, data_only=True)

    print("Creating tables (if they don't exist)...")
    Base.metadata.create_all(bind=engine)

    db: Session = SessionLocal()
    try:
        if clear:
            print("Clearing existing data...")
            for model in [AppState, Stickering, Barcode, EditLog, AssetClassCount, Asset, Class, User]:
                db.query(model).delete()
            db.commit()

        # ── 1. Users (Login ID sheet) ──────────────────────────────────────
        print("Importing users...")
        rows = ws_rows(wb, "Login ID")
        count = 0
        for row in rows[1:]:                         # skip header
            username = str_or_empty(row[0])
            password = str_or_empty(row[1])
            role     = str_or_empty(row[2]) or "Editor"
            if not username or not password:
                continue
            existing = db.query(User).filter(User.username == username).first()
            if existing:
                continue
            db.add(User(username=username, password_hash=hash_password(password), role=role))
            count += 1
        db.commit()
        print(f"  -> {count} users imported")

        # ── 2. Classes (Assign & Lock class sheet) ─────────────────────────
        print("Importing classes...")
        rows = ws_rows(wb, "Assign & Lock class")
        count = 0
        for order, row in enumerate(rows[1:]):
            class_name = str_or_empty(row[0])
            if not class_name:
                continue
            assigned_to       = str_or_empty(row[1])
            locked            = str_or_empty(row[2]).lower() == "yes"
            locked_by         = str_or_empty(row[3])
            lock_timestamp    = str_or_empty(row[4])
            sticker_locked    = str_or_empty(row[5]).lower() == "yes"
            sticker_locked_by = str_or_empty(row[6])
            sticker_timestamp = str_or_empty(row[7])
            existing = db.query(Class).filter(Class.class_name == class_name).first()
            if existing:
                continue
            db.add(Class(
                class_name        = class_name,
                assigned_to       = assigned_to or None,
                locked            = locked,
                locked_by         = locked_by or None,
                lock_timestamp    = lock_timestamp or None,
                sticker_locked    = sticker_locked,
                sticker_locked_by = sticker_locked_by or None,
                sticker_timestamp = sticker_timestamp or None,
                display_order     = order,
            ))
            count += 1
        db.commit()
        print(f"  -> {count} classes imported")

        # ── 3. Assets + class counts (Data sheet) ──────────────────────────
        print("Importing assets and counts...")
        rows = ws_rows(wb, "Data")
        if not rows:
            print("  [skip] Data sheet empty")
        else:
            headers      = rows[0]
            class_names  = [str_or_empty(h) for h in headers[3:]]   # columns D onward

            # Ensure all class columns exist as Class rows
            for i, cn in enumerate(class_names):
                if not cn:
                    continue
                if not db.query(Class).filter(Class.class_name == cn).first():
                    db.add(Class(class_name=cn, display_order=i))
            db.commit()

            # Refresh class lookup
            cls_map = {c.class_name: c for c in db.query(Class).all()}

            asset_count = 0
            count_count = 0
            for row in rows[1:]:
                item_name  = str_or_empty(row[0])
                sl_no_raw  = row[1]
                asset_code = str_or_empty(row[2])
                if not item_name and not asset_code:
                    continue

                sl_no = int(sl_no_raw) if sl_no_raw is not None else None

                asset = db.query(Asset).filter(Asset.asset_code == asset_code).first()
                if not asset:
                    asset = Asset(item_name=item_name, sl_no=sl_no, asset_code=asset_code)
                    db.add(asset)
                    db.flush()           # get asset.id
                    asset_count += 1

                for ci, cn in enumerate(class_names):
                    if not cn or cn not in cls_map:
                        continue
                    raw = row[3 + ci]
                    qty = float(raw) if raw is not None else 0
                    if qty == 0:
                        continue
                    cls_obj = cls_map[cn]
                    exists = db.query(AssetClassCount).filter(
                        AssetClassCount.asset_id == asset.id,
                        AssetClassCount.class_id == cls_obj.id,
                    ).first()
                    if not exists:
                        db.add(AssetClassCount(asset_id=asset.id, class_id=cls_obj.id, count=qty))
                        count_count += 1

            db.commit()
            print(f"  -> {asset_count} assets, {count_count} class-count records imported")

        # ── 4. Barcodes (Bar Code sheet, data starts at row 5) ─────────────
        print("Importing barcodes...")
        rows = ws_rows(wb, "Bar Code")
        count = 0
        for row in rows[4:]:             # index 4 = Excel row 5
            sl_no      = row[0]
            code_text  = str_or_empty(row[1])
            item_name  = str_or_empty(row[2])
            class_name = str_or_empty(row[3])
            remarks    = str_or_empty(row[4])
            if not sl_no and not code_text:
                continue
            db.add(Barcode(
                sl_no      = int(sl_no) if sl_no else None,
                code_text  = code_text,
                item_name  = item_name,
                class_name = class_name,
                remarks    = remarks,
            ))
            count += 1
        db.commit()
        print(f"  -> {count} barcodes imported")

        # ── 5. Stickering ──────────────────────────────────────────────────
        print("Importing stickering records...")
        rows = ws_rows(wb, "Stickering")
        count = 0
        for row in rows[1:]:
            sl_no      = row[0]
            bar_code   = str_or_empty(row[1]).upper()
            item_name  = str_or_empty(row[2])
            class_name = str_or_empty(row[3])
            remarks    = str_or_empty(row[4])
            if not sl_no and not bar_code:
                continue
            existing = db.query(Stickering).filter(Stickering.bar_code == bar_code).first()
            if existing:
                continue
            db.add(Stickering(
                sl_no      = int(sl_no) if sl_no else None,
                bar_code   = bar_code,
                item_name  = item_name,
                class_name = class_name,
                remarks    = remarks,
            ))
            count += 1
        db.commit()
        print(f"  -> {count} stickering records imported")

        # ── 6. Edit Log ────────────────────────────────────────────────────
        print("Importing edit log...")
        rows = ws_rows(wb, "Edit Log")
        count = 0
        for row in rows[1:]:
            username   = str_or_empty(row[0])
            item_name  = str_or_empty(row[1])
            asset_code = str_or_empty(row[2])
            class_name = str_or_empty(row[3])
            qty_raw    = row[4]
            ts_raw     = row[5]
            if not username:
                continue
            qty = float(qty_raw) if qty_raw is not None else 0
            # Try to parse the timestamp; fall back to None (DB uses now())
            created_at = None
            if isinstance(ts_raw, datetime):
                created_at = ts_raw
            db.add(EditLog(
                username   = username,
                item_name  = item_name,
                asset_code = asset_code,
                class_name = class_name,
                qty        = qty,
                created_at = created_at,
            ))
            count += 1
        db.commit()
        print(f"  -> {count} log entries imported")

        # ── Seed lastModified ──────────────────────────────────────────────
        if not db.get(AppState, "lastModified"):
            import time
            db.add(AppState(key="lastModified", value=str(int(time.time() * 1000))))
            db.commit()

        print("\nMigration complete.")

    finally:
        db.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Migrate Excel data to PostgreSQL")
    parser.add_argument(
        "--excel",
        default=os.path.join(os.path.dirname(__file__), "..", "Physical Verification of Fixed Assets.xlsx"),
        help="Path to the Excel file",
    )
    parser.add_argument("--clear", action="store_true", help="Delete existing rows before importing")
    args = parser.parse_args()
    migrate(os.path.abspath(args.excel), args.clear)
