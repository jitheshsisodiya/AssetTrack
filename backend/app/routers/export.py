import io
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from sqlalchemy.orm import Session

from app.auth.jwt import get_current_user
from app.database import get_db
from app.models.models import Asset, AssetClassCount, Barcode, Branch, Class, EditLog, Stickering, User

router = APIRouter(prefix="/api/export", tags=["export"])


def _require_admin(user: User):
    if user.role != "Admin":
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Admin only.")


@router.get("")
def export_data(
    branch_ids: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    _require_admin(current_user)
    iid = current_user.institution_id
    ids = [int(x) for x in branch_ids.split(",") if x.strip().isdigit()] if branch_ids else None

    q = db.query(Class).filter(Class.institution_id == iid)
    if ids:
        q = q.filter(Class.branch_id.in_(ids))
    classes     = q.order_by(Class.display_order).all()
    class_ids   = {c.id for c in classes}
    class_names = [c.class_name for c in classes]

    wb = Workbook()

    ws = wb.active
    ws.title = "Data"
    ws.append(["Item Name", "Sl. No.", "Asset Code"] + class_names)
    asset_ids_q = db.query(AssetClassCount.asset_id).filter(AssetClassCount.class_id.in_(class_ids)).distinct() if class_ids else db.query(Asset.id).filter(Asset.institution_id == iid)
    for asset in db.query(Asset).filter(Asset.institution_id == iid, Asset.id.in_(asset_ids_q)).order_by(Asset.sl_no).all():
        count_map = {acc.class_id: float(acc.count) for acc in asset.counts if acc.class_id in class_ids}
        ws.append([asset.item_name, asset.sl_no, asset.asset_code] + [count_map.get(c.id, 0) or "" for c in classes])

    ws2 = wb.create_sheet("Assign & Lock class")
    ws2.append(["Class Name", "Assigned To", "Locked", "Locked By", "Lock Timestamp", "Sticker Locked", "Sticker Locked By", "Sticker Timestamp", "Branch"])
    branch_map = {b.id: b.name for b in db.query(Branch).filter(Branch.institution_id == iid).all()}
    for c in classes:
        ws2.append([c.class_name, c.assigned_to or "", "Yes" if c.locked else "No", c.locked_by or "", c.lock_timestamp or "",
                    "Yes" if c.sticker_locked else "No", c.sticker_locked_by or "", c.sticker_timestamp or "",
                    branch_map.get(c.branch_id, "") if c.branch_id else ""])

    ws3 = wb.create_sheet("Bar Code")
    ws3.append(["Sl. No.", "Bar Code", "Item Name", "Class Name", "Remarks"])
    bc_q = db.query(Barcode).filter(Barcode.institution_id == iid)
    if ids:
        bc_q = bc_q.filter(Barcode.class_name.in_(class_names))
    for b in bc_q.order_by(Barcode.sl_no).all():
        ws3.append([b.sl_no, b.code_text, b.item_name, b.class_name, b.remarks or ""])

    ws4 = wb.create_sheet("Stickering")
    ws4.append(["Sl. No.", "Bar Code", "Item Name", "Class Name", "Remarks"])
    st_q = db.query(Stickering).filter(Stickering.institution_id == iid)
    if ids:
        st_q = st_q.filter(Stickering.class_name.in_(class_names))
    for s in st_q.order_by(Stickering.sl_no).all():
        ws4.append([s.sl_no, s.bar_code, s.item_name, s.class_name, s.remarks or ""])

    ws5 = wb.create_sheet("Edit Log")
    ws5.append(["Username", "Item Name", "Asset Code", "Class Name", "Qty", "Timestamp"])
    el_q = db.query(EditLog).filter(EditLog.institution_id == iid)
    if ids:
        el_q = el_q.filter(EditLog.class_name.in_(class_names))
    for e in el_q.order_by(EditLog.created_at).all():
        ws5.append([e.username, e.item_name, e.asset_code, e.class_name, float(e.qty) if e.qty else 0, e.created_at.isoformat() if e.created_at else ""])

    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)
    branch_label = "all" if not ids else "-".join(str(i) for i in ids)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="asset_export_{branch_label}.xlsx"'})
